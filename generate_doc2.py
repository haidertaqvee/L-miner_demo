"""
Generator for Deliverable 2: Competition Poster Specification & Blueprint
Table 6 Scoring Rules: Item 2 (Poster, Until 31st Oct - 10 Points + Artistry Crossover)
Generates both .docx and .pdf files.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
)
from reportlab.pdfgen import canvas

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_decorations(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_decorations(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        if self._pageNumber > 1:
            self.drawString(54, 842 - 36, "L-MINER Mk-II | Competition Poster Specification & Blueprint")
            self.drawRightString(595 - 54, 842 - 36, "ISSSP 2026 — Track 3 Finals")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(54, 842 - 42, 595 - 54, 842 - 42)

        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 48, 595 - 54, 48)
        self.drawString(54, 34, "IST Islamabad | Track 3 Finals Scoring Criterion 2 (10 Points)")
        self.drawRightString(595 - 54, 34, f"Page {self._pageNumber} of {page_count}")
        self.restoreState()


def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_docx_callout(doc, text, title=None, border_color="00A8E8", bg_color="F0F9FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    if title:
        run_title = p.add_run(f"{title}\n")
        run_title.bold = True
        run_title.font.name = "Segoe UI"
        run_title.font.size = Pt(10)
        run_title.font.color.rgb = RGBColor(11, 37, 69)
    run_text = p.add_run(text)
    run_text.font.name = "Segoe UI"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ─────────────────────────────────────────────────────────────
# BUILD DOCX: DELIVERABLE 2
# ─────────────────────────────────────────────────────────────
def build_doc2_docx(filepath):
    doc = docx.Document()
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    PRIMARY = RGBColor(11, 37, 69)
    SECONDARY = RGBColor(19, 64, 116)
    ACCENT = RGBColor(0, 168, 232)
    DARK_TEXT = RGBColor(30, 41, 59)
    MUTED_TEXT = RGBColor(100, 116, 139)

    footer = doc.sections[0].footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run("L-MINER Mk-II | Poster Specification (Track 3) · IST Islamabad")
    f_run.font.size = Pt(8)
    f_run.font.color.rgb = MUTED_TEXT

    p_meta = doc.add_paragraph()
    r_tag = p_meta.add_run("ISSSP 2026 — TRACK 3: LUNAR SURFACE ROBOTICS · FINALS SCORING CRITERION 2\n")
    r_tag.font.name = "Segoe UI"; r_tag.font.size = Pt(9); r_tag.font.bold = True; r_tag.font.color.rgb = ACCENT

    r_title = p_meta.add_run("Official Competition Poster Specification & Blueprint")
    r_title.font.name = "Segoe UI"; r_title.font.size = Pt(22); r_title.font.bold = True; r_title.font.color.rgb = PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(
        "Complete Scientific Poster Copy, Technical Schemes, 3-Column Architectural Layout, "
        "and Artistry Design Guidelines (Deadline: 31st Oct — Max Score: 10/10)"
    )
    r_sub.font.name = "Segoe UI"; r_sub.font.size = Pt(11); r_sub.font.color.rgb = SECONDARY

    # Meta Table
    tbl_meta = doc.add_table(rows=2, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Poster Title:", "L-MINER Mk-II: Autonomous Science Payload",
         "Format & Dims:", "ISO A0 Portrait (841 x 1189 mm) / 300 DPI"),
        ("Scoring Focus:", "Table 6: Item 2 (Project Scheme, Background, Methods)",
         "Author Roster:", "Haider Taqvee, Abdul Wasay, Saad, Hajra (IST Islamabad)")
    ]
    for r_idx, (k1, v1, k2, v2) in enumerate(meta_data):
        row = tbl_meta.rows[r_idx]
        c1, c2 = row.cells[0], row.cells[1]
        set_cell_background(c1, "F8FAFC"); set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 80, 80, 100, 100); set_cell_margins(c2, 80, 80, 100, 100)
        p1 = c1.paragraphs[0]; r = p1.add_run(f"{k1} "); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = PRIMARY
        r2 = p1.add_run(v1); r2.font.size = Pt(9); r2.font.color.rgb = DARK_TEXT
        p2 = c2.paragraphs[0]; r = p2.add_run(f"{k2} "); r.bold = True; r.font.size = Pt(9); r.font.color.rgb = PRIMARY
        r2 = p2.add_run(v2); r2.font.size = Pt(9); r2.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── Section 1 ──
    h1 = doc.add_heading("1. Poster Design Specifications & Scoring Mandate", level=1)
    h1.paragraph_format.space_before = Pt(14); h1.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.add_run(
        "According to Table 6 (Finals Scoring Rules of Track 3), the Poster deliverable (due 31st October) carries 10 points. "
        "The evaluators score the poster on: 'Project scheme display, Include Team Information, Background Introduction, "
        "Research Content And Methods, etc.' Furthermore, it directly influences the 5-point 'Artistry' score in the Finals."
    )

    specs = [
        ("Physical Dimensions: ", "Standard ISO A0 Portrait (841 mm width x 1189 mm height; 33.1 x 46.8 inches)."),
        ("Digital Canvas: ", "9933 x 14043 pixels at 300 DPI (CMYK Color Space for professional high-contrast printing)."),
        ("Grid System: ", "3-Column Symmetric Grid with 30 mm margins and 20 mm inter-column gutters."),
        ("Visual Flow: ", "Top Header -> Top-to-Bottom Column 1 (Context & Autonomy) -> Column 2 (Sensors & ML) -> Column 3 (Simulation HUD & ISRU Yield).")
    ]
    for stitle, sdesc in specs:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_after = Pt(3); bp.paragraph_format.line_spacing = 1.15
        r_bt = bp.add_run(stitle); r_bt.bold = True; r_bt.font.color.rgb = PRIMARY
        r_bd = bp.add_run(sdesc)

    add_docx_callout(
        doc,
        "Poster Evaluation Strategy: Judges evaluate scientific clarity, novelty of the technical scheme, "
        "and aesthetic readability from a distance of 1.5 to 2.0 meters. Bullet points are concise, diagrams are high-resolution, "
        "and equations are explicitly annotated.",
        title="Judge Scoring Alignment"
    )

    # ── Section 2: Header Section Copy ──
    h2 = doc.add_heading("2. Poster Header & Team Information Block", level=1)
    h2.paragraph_format.space_before = Pt(14); h2.paragraph_format.space_after = Pt(6)

    header_block = (
        "POSTER TITLE (96 pt Bold):\n"
        "L-MINER Mk-II: Autonomous Multi-Tier Payload for Subsurface Lunar Volatile Prospecting and Mineral Classification\n\n"
        "SUBTITLE (44 pt Medium):\n"
        "Dual-Modality Subsurface GPR and VNIR-SWIR Spectrometry with Edge Machine Learning for Lunar South Pole Cold Traps\n\n"
        "AUTHOR LIST (32 pt Semi-Bold):\n"
        "Haider Taqvee¹, Abdul Wasay¹, Saad¹, Hajra¹\n\n"
        "AFFILIATION (24 pt Regular):\n"
        "¹ Department of Space Science & Aeronautics, Institute of Space Technology (IST), Islamabad, Pakistan\n"
        "ISSSP 2026 — Track 3: Lunar Surface Robotics and ISRU Exploration | Project ID: LMINER-IST-2026"
    )
    add_docx_callout(doc, header_block, title="Header Panel Copy (Exact Print Text)", border_color="0B2545", bg_color="F8FAFC")

    # ── Section 3: Panel Copy ──
    h3 = doc.add_heading("3. Print-Ready Text Copy Across All Poster Panels", level=1)
    h3.paragraph_format.space_before = Pt(14); h3.paragraph_format.space_after = Pt(6)

    panels = [
        ("Panel 1: Background Introduction & Scientific Problem",
         "The Lunar South Pole features Permanently Shadowed Regions (PSRs) where temperatures remain below 40 Kelvin. "
         "Orbital datasets (Chandrayaan-1 M³, LRO LAMP) confirm surface frost, but cannot resolve subsurface stratigraphy or volume.\n\n"
         "The Ground-Truth Gap: Optical cameras are blinded by darkness and only sample the top 10 microns of dust. "
         "Earth teleoperation is impossible due to communication shadows and 3-second roundtrip delays.\n\n"
         "Core Mission Objective: Develop an autonomous prospecting payload combining subsurface radar sounding (0–3 m) "
         "with active optical spectroscopy (0.5–2.5 μm) and edge machine learning to map water ice and titanium-rich minerals in real time."),

        ("Panel 2: Hierarchical Autonomy Architecture (L0 / L1 / L2)",
         "L-MINER payload architecture implements a 3-tier decoupled autonomy framework guaranteeing mission safety and battery preservation:\n\n"
         "• L0 Safety Reflex (10 Hz): Real-time obstacle bypass using HazCam point-clouds. Blocks drive into slopes >15° or boulders >0.75 hazard score.\n\n"
         "• L1 Resource & Thermal Manager (1 Hz): Monitors battery state of charge (SoC) and electronics temperatures. Automatically pauses scans at SoC < 30% or T > 50°C and initiates solar panel recharge cycles.\n\n"
         "• L2 Science Prospecting (On Demand): Executes multi-physics sensor sounding, edge feature engineering, and Random Forest classification upon reaching designated science coordinates."),

        ("Panel 3: Multi-Physics Sensor Simulation & Physical Modeling",
         "Synthetic sensor modeling grounded in Apollo 11/12 sample returns and USGS splib07 spectral libraries:\n\n"
         "• Ground Penetrating Radar (500 MHz): Synthesizes 128 x 256 B-scan profiles across 0–3.0 m depth. "
         "Reflects via dielectric permittivity contrast: Water Ice (A=0.85, width 3), Pyroxene (A=0.65), Ilmenite (A=0.55), Regolith (A=0.30).\n\n"
         "• VNIR-SWIR Spectrometer (0.5–2.5 μm, 200 channels): Active illuminator enables sensing inside dark craters. "
         "Captures diagnostic absorption bands: Water Ice (1.25, 1.5, 2.0 μm), Pyroxene (1.0 & 2.0 μm Fe²⁺ crystal field bands), "
         "Ilmenite (Fe-Ti charge transfer flattening, R ~ 0.07), and space-weathered regolith red slope."),

        ("Panel 4: Edge Feature Extraction & Machine Learning Classification",
         "Edge Data Reduction: Reduces 32,968 raw floats into a 10-dimensional feature vector (3,300x compression):\n"
         "  [0–3] GPR Features: Max amplitude, mean energy, layer standard deviation, peak echo depth.\n"
         "  [4–9] Spectral Features: Mean albedo, R(1.0 μm), R(1.5 μm), R(2.0 μm), spectral slope, 1.5 μm band depth.\n\n"
         "Random Forest Classifier (100 Trees, Gini Impurity):\n"
         "• 95.8% Overall Multiclass Test Accuracy (Water Ice: 99.2%, Ilmenite: 96.5%, Pyroxene: 94.8%, Regolith: 92.5%).\n"
         "• Sub-15 ms Inference Latency on radiation-tolerant ARM Cortex-R5 / RTG4 FPGA embedded processors."),

        ("Panel 5: Interactive 3D Digital Twin & Mission Control HUD",
         "Full cyber-physical demonstration environment built with Streamlit and Plotly WebGL:\n\n"
         "• 3D Lunar Polar Digital Twin: High-density 96 x 96 crater terrain mesh with 3° low-angle polar solar illumination, "
         "6-wheel rocker-bogie mobile host platform, terrain-following pitch/roll adaptation, and subsurface deposit spheres sized by ML confidence.\n\n"
         "• Aerospace Mission HUD: Real-time telemetry cards, dual GPR B-scan + A-scan energy plots, "
         "USGS laboratory benchmark overlay curves, and autonomous survey sequence controls."),

        ("Panel 6: Quantitative ISRU Yields & Mission Impact",
         "In-Situ Resource Utilization (ISRU) Potential:\n"
         "• In a 6 x 6 grid survey (60 x 60 m zone), 7 water ice deposits were localized at average 0.9 m depth.\n"
         "• Yield Calculation: 7 cells x (100 m² x 1.0 m thickness x 1500 kg/m³ density x 5.5 wt% ice) = ~57.8 Tonnes H₂O.\n"
         "• Enables production of ~51.4 Tonnes of Liquid Oxygen (LOX) and ~6.4 Tonnes of Liquid Hydrogen (LH2) propellant, "
         "sufficient to refuel two lunar ascent vehicles for return to Earth orbit.")
    ]

    for p_title, p_content in panels:
        doc.add_heading(p_title, level=2)
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.add_run(p_content)

    # ── Section 4: Visual Layout Wireframe ──
    h4 = doc.add_heading("4. Poster Layout Wireframe & Graphic Design Blueprint", level=1)
    h4.paragraph_format.space_before = Pt(14); h4.paragraph_format.space_after = Pt(6)

    tbl_layout = doc.add_table(rows=5, cols=3)
    tbl_layout.alignment = WD_TABLE_ALIGNMENT.CENTER
    l_headers = ["Column 1 (Width: 247 mm)", "Column 2 (Width: 247 mm)", "Column 3 (Width: 247 mm)"]
    for i, h in enumerate(l_headers):
        cell = tbl_layout.cell(0, i)
        set_cell_background(cell, "0B2545")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]; r = p.add_run(h); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(8.5)

    layout_blocks = [
        ("SECTION 1: MISSION DRIVERS\n• Lunar South Pole Cold Traps\n• Ground Truth Gap diagram\n• High-contrast crater photo",
         "SECTION 3: SENSORS & PHYSICS\n• 500 MHz GPR underbelly diagram\n• VNIR-SWIR Spectrometer arm\n• Dielectric vs Reflectance curves",
         "SECTION 5: 3D DIGITAL TWIN\n• 3D Payload & Host Platform on polar terrain\n• Subterranean deposit ellipsoids\n• GPR sounding radar cone render"),
        ("SECTION 2: HIERARCHICAL AUTONOMY\n• L0 Safety Reflex state diagram\n• L1 Battery & Thermal thresholds\n• L2 Science trigger logic",
         "SECTION 4: MACHINE LEARNING\n• 10-D Feature extraction flowchart\n• Random Forest architecture\n• Confusion Matrix (95.8% Acc)",
         "SECTION 6: ISRU VOLATILE MAP\n• 6x6 Survey grid resource map\n• ~57.8 Tonnes Water Ice calculation\n• LOX/LH2 Propellant production chart"),
        ("FIGURE 1:\nPayload physical layout & host platform integration callouts.",
         "FIGURE 2:\nGPR B-scan paired with A-scan energy profile & USGS comparison.",
         "FIGURE 3:\nLive Streamlit Mission HUD screenshot with real-time telemetry."),
        ("PANEL FOOTER 1:\nFunding & IST Department Acknowledgments.",
         "PANEL FOOTER 2:\nKey References (Apollo, USGS splib07, Chandrayaan-1 M³).",
         "PANEL FOOTER 3:\nQR Code linking to live interactive simulation.")
    ]
    for row_idx, row_vals in enumerate(layout_blocks, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_vals):
            cell = tbl_layout.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 90, 90)
            p = cell.paragraphs[0]; p.paragraph_format.line_spacing = 1.1
            r = p.add_run(val); r.font.size = Pt(8); r.font.color.rgb = DARK_TEXT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ── Section 5: Artistry Palette ──
    h5 = doc.add_heading("5. Color Palette & Typography Specifications (Artistry - 5 Pts)", level=1)
    h5.paragraph_format.space_before = Pt(14); h5.paragraph_format.space_after = Pt(6)

    tbl_pal = doc.add_table(rows=6, cols=4)
    tbl_pal.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_headers = ["Color Role", "Hex Code", "RGB Values", "Application on Poster"]
    for i, h in enumerate(p_headers):
        cell = tbl_pal.cell(0, i)
        set_cell_background(cell, "134074")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]; r = p.add_run(h); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(8.5)

    pal_data = [
        ("Deep Space Black", "#030712", "RGB(3, 7, 18)", "Poster background, header background, terminal blocks"),
        ("Cyan Telemetry Accent", "#00E5FF", "RGB(0, 229, 255)", "Water ice features, radar cones, section titles, key metrics"),
        ("Safety Green", "#00E676", "RGB(0, 230, 118)", "L0 nominal status, Pyroxene mineral indicators, positive callouts"),
        ("Titanium Amber", "#FF9100", "RGB(255, 145, 0)", "Ilmenite signatures, L1 power warnings, dielectric reflectors"),
        ("Hazard Red", "#FF1744", "RGB(255, 23, 68)", "Obstacle pillars, collision warning zones, low-power alerts")
    ]
    for row_idx, row_vals in enumerate(pal_data, start=1):
        bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_vals):
            cell = tbl_pal.cell(row_idx, col_idx)
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 90, 90)
            p = cell.paragraphs[0]; p.paragraph_format.line_spacing = 1.1
            r = p.add_run(val); r.font.size = Pt(8); r.font.color.rgb = DARK_TEXT

    doc.save(filepath)
    print(f"[SUCCESS] Saved DOCX: {filepath}")


# ─────────────────────────────────────────────────────────────
# BUILD PDF: DELIVERABLE 2
# ─────────────────────────────────────────────────────────────
def build_doc2_pdf(filepath):
    doc = SimpleDocTemplate(
        filepath,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )
    styles = getSampleStyleSheet()

    style_meta = ParagraphStyle('MetaTag', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=colors.HexColor("#00A8E8"), spaceAfter=4)
    style_title = ParagraphStyle('DocTitle', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=20, leading=24, textColor=colors.HexColor("#0B2545"), spaceAfter=6)
    style_subtitle = ParagraphStyle('DocSub', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=14, textColor=colors.HexColor("#134074"), spaceAfter=12)
    style_h1 = ParagraphStyle('Heading1_Custom', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=12, leading=15, textColor=colors.HexColor("#0B2545"), spaceBefore=10, spaceAfter=5, keepWithNext=True)
    style_h2 = ParagraphStyle('Heading2_Custom', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=9.5, leading=12, textColor=colors.HexColor("#134074"), spaceBefore=7, spaceAfter=3, keepWithNext=True)
    style_body = ParagraphStyle('Body_Custom', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=11, textColor=colors.HexColor("#1E293B"), spaceAfter=4)
    style_bullet = ParagraphStyle('Bullet_Custom', parent=styles['Normal'], fontName='Helvetica', fontSize=8, leading=11, textColor=colors.HexColor("#1E293B"), leftIndent=12, firstLineIndent=-8, spaceAfter=2)
    style_callout = ParagraphStyle('Callout_Text', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=10.5, textColor=colors.HexColor("#1E293B"))
    style_tbl_cell = ParagraphStyle('Tbl_Cell', parent=styles['Normal'], fontName='Helvetica', fontSize=7, leading=9.5, textColor=colors.HexColor("#1E293B"))
    style_tbl_hdr = ParagraphStyle('Tbl_Hdr', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=10, textColor=colors.white)

    elements = []

    elements.append(Paragraph("ISSSP 2026 — TRACK 3: LUNAR SURFACE ROBOTICS · FINALS SCORING CRITERION 2", style_meta))
    elements.append(Paragraph("Official Competition Poster Specification & Blueprint", style_title))
    elements.append(Paragraph(
        "Complete Scientific Poster Copy, Technical Schemes, 3-Column Architectural Layout, "
        "and Artistry Design Guidelines (Deadline: 31st Oct — Max Score: 10/10)",
        style_subtitle
    ))

    # Meta Table
    meta_data = [
        [Paragraph("<b>Poster Title:</b> L-MINER Mk-II Autonomous Science Payload", style_tbl_cell), Paragraph("<b>Format:</b> ISO A0 Portrait (841 x 1189 mm) / 300 DPI", style_tbl_cell)],
        [Paragraph("<b>Scoring Item:</b> Item 2: Poster (10 Pts)", style_tbl_cell), Paragraph("<b>Authors:</b> Haider, Wasay, Saad, Hajra (IST)", style_tbl_cell)]
    ]
    t_meta = Table(meta_data, colWidths=[240, 247])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F8FAFC")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 6),
        ('RIGHTPADDING', (0,0), (-1,-1), 6),
    ]))
    elements.append(t_meta)
    elements.append(Spacer(1, 8))

    # Section 1
    elements.append(Paragraph("1. Poster Design Specifications & Scoring Mandate", style_h1))
    elements.append(Paragraph(
        "According to Table 6 (Finals Scoring Rules of Track 3), the Poster deliverable (due 31st October) carries 10 points. "
        "Evaluators score on: 'Project scheme display, Include Team Information, Background Introduction, "
        "Research Content And Methods, etc.' It also directly anchors the 5-point 'Artistry' score in the Finals.",
        style_body
    ))
    elements.append(Paragraph("• <b>Physical Dimensions:</b> Standard ISO A0 Portrait (841 mm width x 1189 mm height; 33.1 x 46.8 inches).", style_bullet))
    elements.append(Paragraph("• <b>Digital Canvas:</b> 9933 x 14043 pixels at 300 DPI (CMYK Color Space for professional high-contrast printing).", style_bullet))
    elements.append(Paragraph("• <b>Grid System:</b> 3-Column Symmetric Grid with 30 mm outer margins and 20 mm inter-column gutters.", style_bullet))
    elements.append(Paragraph("• <b>Visual Flow:</b> Header -> Column 1 (Context & Autonomy) -> Column 2 (Sensors & ML) -> Column 3 (HUD & ISRU).", style_bullet))

    # Section 2: Header Copy Callout
    elements.append(Paragraph("2. Poster Header & Team Information Block", style_h1))
    h_text = (
        "<b>POSTER TITLE (96 pt Bold):</b> L-MINER Mk-II: Autonomous Multi-Tier Payload for Subsurface Lunar Volatile Prospecting and Mineral Classification<br/>"
        "<b>SUBTITLE (44 pt Medium):</b> Dual-Modality Subsurface GPR and VNIR-SWIR Spectrometry with Edge Machine Learning for Lunar South Pole Cold Traps<br/>"
        "<b>AUTHORS (32 pt):</b> Haider Taqvee¹, Abdul Wasay¹, Saad¹, Hajra¹<br/>"
        "<b>AFFILIATION (24 pt):</b> ¹ Department of Space Science & Aeronautics, Institute of Space Technology (IST), Islamabad, Pakistan<br/>"
        "<i>ISSSP 2026 — Track 3: Lunar Surface Robotics and ISRU Exploration | Project ID: LMINER-IST-2026</i>"
    )
    t_h = Table([[Paragraph(h_text, style_callout)]], colWidths=[487])
    t_h.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F0F9FF")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#BAE6FD")),
        ('LINELEFT', (0,0), (0,-1), 3, colors.HexColor("#00A8E8")),
        ('TOPPADDING', (0,0), (-1,-1), 5), ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 8), ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    elements.append(t_h)
    elements.append(Spacer(1, 8))

    # Section 3: Panel Copy
    elements.append(Paragraph("3. Print-Ready Text Copy Across All Poster Panels", style_h1))
    
    panel_texts = [
        ("Panel 1: Background & Scientific Problem",
         "The Lunar South Pole features Permanently Shadowed Regions (PSRs) where temperatures remain below 40 K. "
         "Orbital datasets confirm surface frost, but cannot resolve subsurface stratigraphy or volume.<br/>"
         "<b>The Ground-Truth Gap:</b> Optical cameras only sample the top 10 microns of dust. "
         "Earth teleoperation is impossible due to communication shadows and 3-second roundtrip delays.<br/>"
         "<b>Objective:</b> Autonomous rover fusing subsurface radar sounding (0–3 m) with optical spectroscopy (0.5–2.5 μm) and edge ML."),

        ("Panel 2: Hierarchical Autonomy (L0 / L1 / L2)",
         "<b>• L0 Safety Reflex (10 Hz):</b> Real-time obstacle bypass using HazCam point-clouds. Blocks drive into slopes >15° or hazard score >0.75.<br/>"
         "<b>• L1 Resource Manager (1 Hz):</b> Monitors battery SoC and electronics temperatures. Automatically pauses scans at SoC < 30% or T > 50°C for solar recharge.<br/>"
         "<b>• L2 Science Prospecting (On Demand):</b> Executes multi-physics sensor sounding, feature extraction, and ML classification at target coordinates."),

        ("Panel 3: Multi-Physics Sensor Modeling",
         "Synthetic sensor modeling grounded in Apollo sample returns and USGS splib07 library:<br/>"
         "<b>• 500 MHz GPR:</b> 128 x 256 B-scan profiles across 0–3 m depth via dielectric permittivity contrast (Ice: A=0.85, Ilmenite: A=0.55, Pyroxene: A=0.65).<br/>"
         "<b>• VNIR-SWIR Spectrometer (0.5–2.5 μm):</b> Active illumination in dark PSRs. Captures diagnostic bands: Water Ice (1.25, 1.5, 2.0 μm), Pyroxene (1.0 & 2.0 μm), Ilmenite (flat R~0.07)."),

        ("Panel 4: Edge Feature Extraction & ML Engine",
         "<b>• 3,300x Edge Data Reduction:</b> Compresses 32,968 raw floats into a 10-D feature vector (4 GPR features + 6 spectral features).<br/>"
         "<b>• Random Forest Classifier (100 Trees):</b> 95.8% overall multiclass accuracy (Ice: 99.2%, Ilm: 96.5%, Pyr: 94.8%, Reg: 92.5%).<br/>"
         "<b>• Sub-15 ms Inference Latency</b> on radiation-hardened ARM Cortex-R5 / RTG4 FPGA processors."),

        ("Panel 5: Interactive 3D Digital Twin & Mission HUD",
         "<b>• 3D Digital Twin:</b> High-density 96 x 96 crater terrain mesh with 3° polar lighting, 6-wheel rocker-bogie rover, terrain slope adaptation, and subterranean deposit ellipsoids sized by ML confidence.<br/>"
         "<b>• Streamlit Mission Control:</b> Real-time telemetry, dual GPR B-scan + A-scan energy plots, USGS benchmark overlay, and survey controls."),

        ("Panel 6: Quantitative ISRU Yields & Mission Impact",
         "<b>• Survey Yield:</b> 7 water ice deposits localized at avg 0.9 m depth across 60 x 60 m survey zone.<br/>"
         "<b>• Resource Potential:</b> ~57.8 Tonnes H₂O mapped (1500 kg/m³ density, 5.5 wt% ice).<br/>"
         "<b>• Propellant Output:</b> ~51.4 Tonnes LOX + ~6.4 Tonnes LH2, capable of refueling two lunar ascent return vehicles.")
    ]

    for p_title, p_body in panel_texts:
        elements.append(Paragraph(f"<b>{p_title}</b>", style_h2))
        elements.append(Paragraph(p_body, style_body))

    # Section 4: Layout Table
    elements.append(Spacer(1, 4))
    elements.append(Paragraph("4. Poster Layout Wireframe & Graphic Design Blueprint", style_h1))
    layout_table_data = [
        [Paragraph("Column 1 (247 mm)", style_tbl_hdr), Paragraph("Column 2 (247 mm)", style_tbl_hdr), Paragraph("Column 3 (247 mm)", style_tbl_hdr)],
        [Paragraph("<b>Panel 1: Context & Gap</b><br/>• Lunar South Pole Cold Traps<br/>• In-situ prospecting challenge", style_tbl_cell),
         Paragraph("<b>Panel 3: Sensor Physics</b><br/>• 500 MHz GPR profile<br/>• VNIR-SWIR absorption curves", style_tbl_cell),
         Paragraph("<b>Panel 5: 3D Digital Twin</b><br/>• Rocker-bogie 3D rover mesh<br/>• Subsurface deposit spheres", style_tbl_cell)],
        [Paragraph("<b>Panel 2: Autonomy</b><br/>• L0 Reflex, L1 Resource, L2 ML<br/>• Flowchart of state transitions", style_tbl_cell),
         Paragraph("<b>Panel 4: Machine Learning</b><br/>• 10-D Feature extraction<br/>• Confusion matrix & accuracy", style_tbl_cell),
         Paragraph("<b>Panel 6: ISRU Potential</b><br/>• 57.8 Tonnes H₂O calculation<br/>• LOX/LH2 propellant yield", style_tbl_cell)]
    ]
    t_layout = Table(layout_table_data, colWidths=[162, 162, 163])
    t_layout.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#0B2545")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.HexColor("#F8FAFC"), colors.white]),
        ('TOPPADDING', (0,0), (-1,-1), 4), ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 5), ('RIGHTPADDING', (0,0), (-1,-1), 5),
    ]))
    elements.append(t_layout)

    doc.build(elements, canvasmaker=NumberedCanvas)
    print(f"[SUCCESS] Saved PDF: {filepath}")


if __name__ == "__main__":
    out_dir = r"E:\L-miner directory\Finals_Track3_Deliverables"
    os.makedirs(out_dir, exist_ok=True)
    docx_path = os.path.join(out_dir, "2_Poster_Specification_and_Blueprint.docx")
    pdf_path = os.path.join(out_dir, "2_Poster_Specification_and_Blueprint.pdf")
    build_doc2_docx(docx_path)
    build_doc2_pdf(pdf_path)
